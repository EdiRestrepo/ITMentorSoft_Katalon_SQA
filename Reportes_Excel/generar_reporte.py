"""
Genera el reporte de resultados de pruebas de ITMentorSoft en Excel y HTML.

Fuentes:
  - casos_azure.json: los 43 test cases de Azure DevOps (catálogo leído por MCP en modo lectura), su HU,
    responsable, tipo (Automatizado / Automatizado parcial / No automatizable) y el caso Katalon que los cubre.
  - hallazgos.json: defectos y observaciones encontrados al ejecutar las pruebas.
  - ../Reports/**/JUnit_Report.xml y execution0.log: ejecuciones de las suites de Katalon.
  - ../Evidencias/<TC>/bitacora.txt: bitácora de cada caso (carpeta local no versionada; si falta, se omite).

Uso (desde esta carpeta):  python generar_reporte.py
Salida:
  - ../index.html                              reporte web (se puede publicar con GitHub Pages)
  - Reporte_Resultados_<n>.xlsx                consolidado (n = siguiente número libre; conserva los anteriores)
  - Por_prueba/<caso Katalon>.xlsx             reporte de cada prueba (TC01 … TC20)
"""
import html
import json
import re
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
PROYECTO = AQUI.parent
AZURE = 'https://dev.azure.com/especializacion-itm/Trabajo%20de%20Grado%20-%20Tutor%20inteligente/_workitems/edit/'
APP = 'https://itmentorsoft.com'

RESULTADOS = ['Pasó', 'Falló', 'Error', 'No ejecutado', 'No automatizable']
ESTADO_KATALON = {'PASSED': 'Pasó', 'FAILED': 'Falló', 'ERROR': 'Error', 'INCOMPLETE': 'Error', 'SKIPPED': 'No ejecutado'}
CLASE = {'Pasó': 'ok', 'Falló': 'ko', 'Error': 'err', 'No ejecutado': 'ne', 'No automatizable': 'na'}
COLOR = {'Pasó': 'D9F2E3', 'Falló': 'FAD7D4', 'Error': 'F8E1C4', 'No ejecutado': 'EDEFF2', 'No automatizable': 'E8E0F5'}
TIPOS = ['Automatizado', 'Automatizado parcial', 'No automatizable']


# ---------------------------------------------------------------- datos
def cargar_json(nombre):
    return json.loads((AQUI / nombre).read_text(encoding='utf-8'))


def casos_katalon():
    """'TC04' -> nombre del test case en Katalon (por el prefijo TCnn_ del archivo .tc)."""
    mapa = {}
    for tc in sorted((PROYECTO / 'Test Cases').glob('TC*.tc')):
        m = re.match(r'(TC\d+)_', tc.name)
        if m:
            mapa[m.group(1)] = tc.stem
    return mapa


def a_local(texto):
    """Fecha ISO de Katalon (UTC) -> datetime local sin zona."""
    valor = datetime.fromisoformat(texto.replace('Z', '+00:00'))
    return (valor if valor.tzinfo else valor.replace(tzinfo=timezone.utc)).astimezone().replace(tzinfo=None)


def campo_log(bloque, nombre):
    m = re.search(r'<%s>(.*?)</%s>' % (nombre, nombre), bloque, re.S)
    texto, anterior = (m.group(1) if m else ''), None
    while texto != anterior:  # Katalon escapa algunos mensajes dos veces
        anterior, texto = texto, html.unescape(texto)
    return texto.strip()


def mensajes_log(carpeta, caso):
    """Advertencias y fallos del caso registrados en execution0.log."""
    log = Path(carpeta) / 'execution0.log'
    if not log.exists():
        return []
    marca, salida = f'[{caso}]', []
    for bloque in re.findall(r'<record>(.*?)</record>', log.read_text(encoding='utf-8', errors='replace'), re.S):
        nivel, mensaje = campo_log(bloque, 'level'), campo_log(bloque, 'message')
        if nivel in ('WARNING', 'FAILED', 'ERROR') and marca in mensaje:
            salida.append((nivel, mensaje.split('\n')[0]))
    return list(dict.fromkeys(salida))


def leer_ejecuciones(katalon):
    """Una entrada por test case ejecutado en cada JUnit_Report.xml de Reports/."""
    ejecuciones = []
    for xml in (PROYECTO / 'Reports').rglob('JUnit_Report.xml'):
        try:
            raiz = ET.parse(xml).getroot()
        except ET.ParseError:
            continue
        for suite in raiz.iter('testsuite'):
            props = {p.get('name'): p.get('value') for p in suite.iter('property')}
            for tc in suite.iter('testcase'):
                m = re.search(r'(TC\d+)_', tc.get('name', ''))
                if not m or m.group(1) not in katalon:
                    continue
                falla = tc.find('failure') if tc.find('failure') is not None else tc.find('error')
                inicio = a_local(suite.get('timestamp'))
                ejecuciones.append({
                    'caso': m.group(1), 'nombre': katalon[m.group(1)], 'suite': suite.get('name'),
                    'fecha': inicio, 'duracion': round(float(tc.get('time') or 0), 1),
                    'resultado': ESTADO_KATALON.get((tc.get('status') or '').upper(), 'Error'),
                    'mensaje': (falla.get('message') or '').split('\n')[0][:600] if falla is not None else '',
                    'reporte': xml.parent.relative_to(PROYECTO).as_posix(),
                    'log': props.get('logFolder', str(xml.parent)),
                    'capturas': [a.strip() for a in (props.get('attachments') or '').split(',') if a.strip()],
                })
    return sorted(ejecuciones, key=lambda e: e['fecha'])


def bitacora(caso, desde, hasta):
    archivo = PROYECTO / 'Evidencias' / caso / 'bitacora.txt'
    if not archivo.exists():
        return []
    lineas = []
    for linea in archivo.read_text(encoding='utf-8', errors='replace').splitlines():
        m = re.match(r'(\d{8}_\d{6})_\d{3} \| (.*)', linea)
        if m and desde <= datetime.strptime(m.group(1), '%Y%m%d_%H%M%S') <= hasta and not m.group(2).startswith('Captura:'):
            lineas.append((datetime.strptime(m.group(1), '%Y%m%d_%H%M%S'), m.group(2)))
    return lineas


def evidencia_versionada(caso, ruta):
    """Ruta relativa de la captura copiada en Documentacion/Evidencias_Ejecucion (si existe) o la original."""
    nombre = Path(ruta).name
    copia = PROYECTO / 'Documentacion' / 'Evidencias_Ejecucion' / caso / nombre
    return (copia if copia.exists() else Path(ruta)).as_posix().replace(PROYECTO.as_posix() + '/', '')


def preparar():
    datos, hz = cargar_json('casos_azure.json'), cargar_json('hallazgos.json')
    katalon = casos_katalon()
    ejecuciones = leer_ejecuciones(katalon)
    ultima = {}
    for e in ejecuciones:
        ultima[e['caso']] = e
    hallazgos_de = {}
    for h in hz['hallazgos']:
        for c in h['casos']:
            hallazgos_de.setdefault(c, []).append(h['id'])

    filas = []
    for c in datos['casos']:
        e = ultima.get(c['katalon']) if c['katalon'] else None
        resultado = 'No automatizable' if c['tipo'] == 'No automatizable' else (e['resultado'] if e else 'No ejecutado')
        filas.append(dict(c, katalon_nombre=katalon.get(c['katalon'], ''), suite=e['suite'] if e else '',
                          resultado=resultado, duracion=e['duracion'] if e else None,
                          fecha=e['fecha'].strftime('%Y-%m-%d %H:%M:%S') if e else '',
                          fuente=f"Katalon ({e['suite']})" if e else ('Clasificación del inventario' if c['tipo'] == 'No automatizable' else ''),
                          mensaje=e['mensaje'] if e else '', hallazgos=', '.join(hallazgos_de.get(c['id'], []))))

    pruebas = []
    for caso, nombre in katalon.items():
        historial = [e for e in ejecuciones if e['caso'] == caso]
        u = historial[-1] if historial else None
        azure = [f for f in filas if f['katalon'] == caso]
        pasos, msgs, capturas = [], [], []
        if u:
            from datetime import timedelta
            pasos = bitacora(caso, u['fecha'] - timedelta(seconds=5), u['fecha'] + timedelta(seconds=u['duracion'] + 60))
            msgs = mensajes_log(u['log'], caso)
            capturas = [evidencia_versionada(caso, r) for r in u['capturas']]
        pruebas.append({'caso': caso, 'nombre': nombre, 'azure': azure, 'ultima': u, 'historial': historial,
                        'pasos': pasos, 'mensajes': msgs, 'capturas': capturas,
                        'hallazgos': sorted({h for f in azure for h in f['hallazgos'].split(', ') if h})})
    return datos, hz, filas, pruebas, ejecuciones


# ---------------------------------------------------------------- Excel
F = 'Arial'
BASE, NEGRITA = Font(name=F, size=10), Font(name=F, size=10, bold=True)
CAB_FONT, CAB_FILL = Font(name=F, size=10, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
LINEA = Side(style='thin', color='BFBFBF')
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
AJUSTE = Alignment(wrap_text=True, vertical='top')
CENTRO = Alignment(horizontal='center', vertical='top', wrap_text=True)
ENLACE = Font(name=F, size=10, color='0563C1', underline='single')


def cabecera(ws, fila, textos, anchos):
    for i, (t, a) in enumerate(zip(textos, anchos), 1):
        c = ws.cell(fila, i, t)
        c.font, c.fill, c.alignment, c.border = CAB_FONT, CAB_FILL, CENTRO, BORDE
        ws.column_dimensions[get_column_letter(i)].width = a


def colorear(ws, rango):
    for r, color in COLOR.items():
        ws.conditional_formatting.add(rango, CellIsRule(operator='equal', formula=[f'"{r}"'], fill=PatternFill('solid', fgColor=color)))


def escribir_fila(ws, n, valores, centrados=()):
    for i, v in enumerate(valores, 1):
        c = ws.cell(n, i, v)
        c.font, c.border = BASE, BORDE
        c.alignment = CENTRO if i in centrados else AJUSTE


def titulo(ws, texto, subtitulo):
    ws['A1'] = texto
    ws['A1'].font = Font(name=F, size=14, bold=True)
    ws['A2'] = subtitulo
    ws['A2'].font = Font(name=F, size=9, italic=True)


def hoja_hallazgos(wb, hz, indice=None):
    hal = wb.create_sheet('Hallazgos', indice) if indice is not None else wb.create_sheet('Hallazgos')
    titulo(hal, 'Hallazgos de las pruebas', f"{hz['metodo']} ({hz['fecha']})")
    cabecera(hal, 4, ['ID', 'Categoría', 'Severidad', 'Hallazgo', 'Detalle', 'Evidencia', 'Casos de Azure', 'Estado', 'Recomendación'],
             [7, 13, 11, 40, 60, 45, 16, 34, 40])
    sev_color = {'Crítica': 'F4B6B0', 'Alta': 'F8D3A8', 'Media': 'FFF2CC', 'Baja': 'E2EFDA'}
    for n, h in enumerate(hz['hallazgos'], 5):
        escribir_fila(hal, n, [h['id'], h['categoria'], h['severidad'], h['titulo'], h['detalle'], h['evidencia'],
                               ', '.join(map(str, h['casos'])), h['estado'], h['recomendacion']], centrados=(1, 2, 3))
        hal.cell(n, 1).font = hal.cell(n, 4).font = NEGRITA
        if h['severidad'] in sev_color:
            hal.cell(n, 3).fill = PatternFill('solid', fgColor=sev_color[h['severidad']])
    hal.freeze_panes = 'E5'


def excel_consolidado(datos, hz, filas, pruebas, ejecuciones, destino):
    wb = Workbook()

    # ---- Detalle (un test case de Azure por fila)
    det = wb.active
    det.title = 'Detalle'
    cabecera(det, 1, ['ID Azure', 'Título', 'HU', 'Responsable', 'Tipo', 'Suite', 'Caso Katalon', 'Resultado', 'Duración (s)',
                      'Fecha ejecución', 'Fuente del resultado', 'Mensaje / observación', 'Nota', 'Hallazgos'],
             [9, 48, 34, 26, 18, 12, 46, 15, 11, 18, 22, 50, 40, 12])
    for n, f in enumerate(filas, 2):
        escribir_fila(det, n, [f['id'], f['titulo'], '\n'.join(f"{h['id']} - {h['titulo']}" for h in f['hu']), f['responsable'], f['tipo'],
                               f['suite'], f['katalon_nombre'], f['resultado'], f['duracion'], f['fecha'], f['fuente'], f['mensaje'],
                               f['nota'], f['hallazgos']], centrados=(1, 5, 8, 9, 10))
        det.cell(n, 1).hyperlink = AZURE + str(f['id'])
        det.cell(n, 1).font = ENLACE
    ultima_det = len(filas) + 1
    det.freeze_panes = 'C2'
    det.auto_filter.ref = f'A1:N{ultima_det}'
    colorear(det, f'H2:H{ultima_det}')

    # ---- HU_Caso (una fila por par HU-caso, para contar por HU)
    huc = wb.create_sheet('HU_Caso')
    cabecera(huc, 1, ['HU', 'Título HU', 'ID caso', 'Resultado'], [8, 60, 10, 18])
    n = 2
    for f in filas:
        for h in f['hu']:
            escribir_fila(huc, n, [h['id'], h['titulo'], f['id'],
                                   f'=INDEX(Detalle!$H$2:$H${ultima_det},MATCH(C{n},Detalle!$A$2:$A${ultima_det},0))'])
            n += 1
    ultima_huc = n - 1

    # ---- Casos Katalon (una fila por prueba automatizada)
    kat = wb.create_sheet('Casos_Katalon')
    cabecera(kat, 1, ['Caso', 'Test case Katalon', 'Casos Azure', 'Resultado', 'Fecha ejecución', 'Duración (s)', 'Ejecuciones',
                      'Mensaje', 'Advertencias', 'Hallazgos', 'Reporte de la prueba'], [8, 52, 22, 14, 18, 11, 11, 50, 60, 12, 34])
    for n, p in enumerate(pruebas, 2):
        u = p['ultima']
        escribir_fila(kat, n, [p['caso'], p['nombre'], ', '.join(f"#{f['id']}" for f in p['azure']), u['resultado'] if u else 'No ejecutado',
                               u['fecha'].strftime('%Y-%m-%d %H:%M:%S') if u else '', u['duracion'] if u else None, len(p['historial']),
                               u['mensaje'] if u else '', '\n'.join(m for nv, m in p['mensajes'] if nv == 'WARNING'),
                               ', '.join(p['hallazgos']), f"Por_prueba/{p['nombre']}.xlsx"], centrados=(1, 4, 5, 6, 7))
        kat.cell(n, 11).hyperlink = f"Por_prueba/{p['nombre']}.xlsx"
        kat.cell(n, 11).font = ENLACE
    kat.freeze_panes = 'C2'
    colorear(kat, f'D2:D{len(pruebas) + 1}')

    # ---- Resumen
    res = wb.create_sheet('Resumen', 0)
    titulo(res, 'Resultados de pruebas - ITMentorSoft',
           f"Generado: {datetime.now():%Y-%m-%d %H:%M}  ·  Casos: {datos['fuente']} ({datos['fecha_lectura']})  ·  "
           f"Ejecuciones de Katalon leídas: {len(ejecuciones)}  ·  Aplicación: {APP}")
    cabecera(res, 4, ['Resultado (casos de Azure)', 'Casos', '% del total'], [34, 10, 12])
    D = f'Detalle!$H$2:$H${ultima_det}'
    for i, r in enumerate(RESULTADOS, 5):
        escribir_fila(res, i, [r, f'=COUNTIF({D},A{i})', f'=IF($B$10=0,0,B{i}/$B$10)'])
    escribir_fila(res, 10, ['Total', '=SUM(B5:B9)', '=IF(B10=0,0,B10/B10)'])
    escribir_fila(res, 11, ['Ejecutados', '=B5+B6+B7', None])
    escribir_fila(res, 12, ['% de aprobación (pasó / ejecutados)', '=IF(B11=0,0,B5/B11)', None])
    for r in range(5, 13):
        res.cell(r, 3).number_format = '0%'
        if r >= 10:
            for c in range(1, 4):
                res.cell(r, c).font = NEGRITA
    res['B12'].number_format = '0%'
    colorear(res, 'A5:A9')

    cabecera(res, 14, ['Tipo de caso', 'Casos', 'Pasó', 'Falló / Error', 'Sin ejecutar'], [34, 10, 12, 14, 14])
    T = f'Detalle!$E$2:$E${ultima_det}'
    for i, t in enumerate(TIPOS, 15):
        escribir_fila(res, i, [t, f'=COUNTIF({T},A{i})', f'=COUNTIFS({T},A{i},{D},"Pasó")',
                               f'=COUNTIFS({T},A{i},{D},"Falló")+COUNTIFS({T},A{i},{D},"Error")', f'=B{i}-C{i}-D{i}'])

    K = f'Casos_Katalon!$D$2:$D${len(pruebas) + 1}'
    cabecera(res, 19, ['Pruebas Katalon', 'Casos', '', '', ''], [34, 10, 12, 14, 14])
    for i, (t, formula) in enumerate([('Pruebas automatizadas', f'=COUNTA(Casos_Katalon!$A$2:$A${len(pruebas) + 1})'),
                                      ('Pasaron', f'=COUNTIF({K},"Pasó")'), ('Fallaron', f'=COUNTIF({K},"Falló")'),
                                      ('Error', f'=COUNTIF({K},"Error")'), ('No ejecutadas', f'=COUNTIF({K},"No ejecutado")')], 20):
        escribir_fila(res, i, [t, formula])

    cabecera(res, 26, ['HU', 'Historia de usuario', 'Casos', 'Pasó', 'Falló / Error', 'Sin ejecutar'], [34, 60, 10, 12, 14, 14])
    H, R = f'HU_Caso!$A$2:$A${ultima_huc}', f'HU_Caso!$D$2:$D${ultima_huc}'
    hus = sorted({(h['id'], h['titulo']) for f in filas for h in f['hu']}, key=lambda x: x[1])
    for i, (hid, htit) in enumerate(hus, 27):
        escribir_fila(res, i, [hid, htit, f'=COUNTIF({H},A{i})', f'=COUNTIFS({H},A{i},{R},"Pasó")',
                               f'=COUNTIFS({H},A{i},{R},"Falló")+COUNTIFS({H},A{i},{R},"Error")', f'=C{i}-D{i}-E{i}'], centrados=(1,))
    res.cell(27 + len(hus) + 1, 1, '"Sin ejecutar" incluye los casos clasificados como No automatizable.').font = Font(name=F, size=9, italic=True)
    res.column_dimensions['B'].width = 60

    hoja_hallazgos(wb, hz, 1)

    # ---- Ejecuciones (historial completo)
    ej = wb.create_sheet('Ejecuciones')
    cabecera(ej, 1, ['Fecha', 'Suite', 'Caso', 'Test case Katalon', 'Resultado', 'Duración (s)', 'Mensaje', 'Reporte'], [18, 12, 8, 52, 13, 11, 60, 48])
    for n, e in enumerate(reversed(ejecuciones), 2):
        escribir_fila(ej, n, [e['fecha'].strftime('%Y-%m-%d %H:%M:%S'), e['suite'], e['caso'], e['nombre'], e['resultado'], e['duracion'],
                              e['mensaje'], e['reporte']], centrados=(1, 3, 5, 6))
    if ejecuciones:
        colorear(ej, f'E2:E{len(ejecuciones) + 1}')
    ej.freeze_panes = 'A2'

    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)


def excel_prueba(p, destino):
    wb = Workbook()
    res = wb.active
    res.title = 'Resumen'
    u = p['ultima']
    titulo(res, f"{p['caso']} - {p['nombre']}", f"Reporte de la prueba · Generado: {datetime.now():%Y-%m-%d %H:%M} · Aplicación: {APP}")
    filas = [('Caso Katalon', f"Test Cases/{p['nombre']}"), ('Resultado de la última ejecución', u['resultado'] if u else 'No ejecutado'),
             ('Fecha de ejecución', u['fecha'].strftime('%Y-%m-%d %H:%M:%S') if u else ''), ('Duración (s)', u['duracion'] if u else None),
             ('Suite', u['suite'] if u else ''), ('Reporte de Katalon', u['reporte'] if u else ''), ('Ejecuciones registradas', len(p['historial'])),
             ('Mensaje de fallo', u['mensaje'] if u else ''), ('Hallazgos relacionados', ', '.join(p['hallazgos']))]
    cabecera(res, 4, ['Campo', 'Valor'], [32, 100])
    for n, (k, v) in enumerate(filas, 5):
        escribir_fila(res, n, [k, v])
        res.cell(n, 1).font = NEGRITA
    colorear(res, 'B6')
    n = 5 + len(filas) + 1
    cabecera(res, n, ['ID Azure', 'Caso de Azure', 'HU', 'Tipo', 'Responsable', 'Nota'], [32, 100, 34, 18, 28, 50])
    for i, f in enumerate(p['azure'], n + 1):
        escribir_fila(res, i, [f['id'], f['titulo'], f['hu'][0]['titulo'], f['tipo'], f['responsable'], f['nota']], centrados=(1,))
        res.cell(i, 1).hyperlink = AZURE + str(f['id'])
        res.cell(i, 1).font = ENLACE

    bit = wb.create_sheet('Bitacora')
    cabecera(bit, 1, ['#', 'Hora', 'Registro'], [6, 12, 140])
    for i, (hora, texto) in enumerate(p['pasos'], 2):
        escribir_fila(bit, i, [i - 1, hora.strftime('%H:%M:%S'), texto], centrados=(1, 2))
    if not p['pasos']:
        bit.cell(2, 3, 'Sin bitácora local (Evidencias/ no se versiona) o sin ejecuciones.').font = Font(name=F, size=10, italic=True)

    msg = wb.create_sheet('Mensajes')
    cabecera(msg, 1, ['Nivel', 'Mensaje'], [12, 140])
    for i, (nivel, texto) in enumerate(p['mensajes'], 2):
        escribir_fila(msg, i, [nivel, texto], centrados=(1,))
        msg.cell(i, 1).fill = PatternFill('solid', fgColor='FAD7D4' if nivel in ('FAILED', 'ERROR') else 'FFF2CC')
    if not p['mensajes']:
        msg.cell(2, 2, 'Sin advertencias ni fallos registrados en la última ejecución.').font = Font(name=F, size=10, italic=True)

    ej = wb.create_sheet('Ejecuciones')
    cabecera(ej, 1, ['Fecha', 'Suite', 'Resultado', 'Duración (s)', 'Mensaje', 'Reporte'], [18, 12, 13, 11, 80, 48])
    for i, e in enumerate(reversed(p['historial']), 2):
        escribir_fila(ej, i, [e['fecha'].strftime('%Y-%m-%d %H:%M:%S'), e['suite'], e['resultado'], e['duracion'], e['mensaje'], e['reporte']],
                      centrados=(1, 3, 4))
    if p['historial']:
        colorear(ej, f'C2:C{len(p["historial"]) + 1}')

    ev = wb.create_sheet('Evidencias')
    cabecera(ev, 1, ['#', 'Captura (ruta relativa al proyecto)'], [6, 120])
    for i, c in enumerate(p['capturas'], 2):
        escribir_fila(ev, i, [i - 1, c], centrados=(1,))
        ev.cell(i, 2).hyperlink = '../../' + c
        ev.cell(i, 2).font = ENLACE
    wb.save(destino)


# ---------------------------------------------------------------- HTML
CSS = """:root{--bg:#f6f7f9;--card:#fff;--tx:#1d2330;--mut:#5f6b7a;--bd:#dde2e8;--ac:#1f4e78;
--ok:#2e8b57;--ko:#c0392b;--err:#d68910;--na:#7d5ba6;--ne:#9aa5b1}
@media (prefers-color-scheme:dark){:root:not([data-theme=light]){--bg:#12161c;--card:#1b2129;--tx:#e6eaf0;--mut:#9aa6b4;--bd:#2c3540;--ac:#6fa8dc}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);font:14px/1.45 system-ui,Segoe UI,Arial,sans-serif}
main{max-width:1200px;margin:0 auto;padding:24px 16px 48px}h1{margin:0 0 4px;font-size:24px}h2{font-size:17px;margin:28px 0 10px}
.meta{color:var(--mut);font-size:12.5px}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:18px 0}
.card{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px}.card b{display:block;font-size:26px}.card span{color:var(--mut);font-size:12.5px}
.bar,.mini{display:flex;height:14px;border-radius:7px;overflow:hidden;background:var(--bd)}.mini{height:8px;min-width:90px}
.ok{background:var(--ok)}.ko{background:var(--ko)}.err{background:var(--err)}.na{background:var(--na)}.ne{background:var(--ne)}
.legend{display:flex;flex-wrap:wrap;gap:14px;list-style:none;padding:0;margin:10px 0 0;font-size:13px}.dot{display:inline-block;width:10px;height:10px;border-radius:50%;margin-right:6px}
.panel{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px;overflow-x:auto}
table{border-collapse:collapse;width:100%;font-size:13px}th,td{text-align:left;padding:7px 8px;border-bottom:1px solid var(--bd);vertical-align:top}
th{color:var(--mut);font-weight:600;font-size:12px;position:sticky;top:0;background:var(--card)}.num{text-align:right;white-space:nowrap}
.sub{color:var(--mut);font-size:11.5px;margin-top:2px}.msg{max-width:380px;font-size:12px;color:var(--mut);word-break:break-word}
.chip{color:#fff;border-radius:999px;padding:2px 9px;font-size:12px;white-space:nowrap}a{color:var(--ac)}
.filters{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:10px}input,select{font:inherit;padding:6px 8px;border:1px solid var(--bd);border-radius:6px;background:var(--card);color:var(--tx)}
input{flex:1;min-width:180px}
.hzgrid{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:12px}
.hzc{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px}.hzc h3{font-size:14.5px;margin:8px 0 6px}.hzc p{margin:6px 0;font-size:13px}
.hzc header{display:flex;align-items:center;gap:8px;flex-wrap:wrap}.sev{color:#fff;border-radius:999px;padding:2px 9px;font-size:11.5px}
.sev.critica{background:#922b21}.sev.alta{background:var(--ko)}.sev.media{background:var(--err)}.sev.baja{background:var(--ne)}
.hz{display:inline-block;border:1px solid var(--bd);border-radius:6px;padding:0 6px;margin:0 4px 2px 0;font-size:11.5px;color:var(--mut)}
.downloads{display:flex;flex-wrap:wrap;gap:8px}.downloads a{background:var(--card);border:1px solid var(--bd);border-radius:8px;padding:8px 12px;text-decoration:none;font-size:13px}
@media (max-width:640px){main{padding:16px}h1{font-size:20px}}"""


def chip(resultado):
    return f'<span class="chip {CLASE.get(resultado, "ne")}">{html.escape(resultado)}</span>'


def barra(conteo, total, clase='bar'):
    return f'<div class="{clase}">' + ''.join(
        f'<span class="{CLASE[r]}" style="flex:{conteo.get(r, 0)}" title="{r}: {conteo.get(r, 0)}"></span>' for r in RESULTADOS) + '</div>'


def html_reporte(datos, hz, filas, pruebas, ejecuciones, excel, destino):
    e = html.escape
    conteo = {r: sum(f['resultado'] == r for f in filas) for r in RESULTADOS}
    ejecutadas = [p for p in pruebas if p['ultima']]
    pasaron = sum(p['ultima']['resultado'] == 'Pasó' for p in ejecutadas)
    automatizados = sum(f['tipo'] != 'No automatizable' for f in filas)
    defectos = sum(h['id'].startswith('D-') for h in hz['hallazgos'])
    tarjetas = [(len(filas), 'Test cases en Azure'), (automatizados, 'Automatizados en Katalon (completo o parcial)'),
                (len(pruebas), 'Pruebas Katalon'), (len(ejecutadas), 'Pruebas ejecutadas'), (pasaron, 'Pasaron'),
                (len(ejecutadas) - pasaron, 'Fallaron / error'),
                (f"{round(100 * pasaron / len(ejecutadas)) if ejecutadas else 0}%", 'Aprobación (pasó / ejecutadas)')]

    hallazgos = ''.join(
        f'<article class="hzc"><header><span class="sev {e(h["severidad"].lower().replace("í", "i"))}">{e(h["severidad"])}</span>'
        f'<b>{e(h["id"])}</b> <span class="sub">{e(h["categoria"])}</span></header><h3>{e(h["titulo"])}</h3><p>{e(h["detalle"])}</p>'
        f'<p class="sub"><b>Estado:</b> {e(h["estado"])}</p><p class="sub"><b>Recomendación:</b> {e(h["recomendacion"])}</p>'
        f'<p class="sub"><b>Evidencia:</b> {e(h["evidencia"])}</p><p class="sub"><b>Casos de Azure:</b> {", ".join(map(str, h["casos"]))}</p></article>'
        for h in hz['hallazgos'])

    filas_kat = ''.join(
        f'<tr><td class="num">{e(p["caso"])}</td><td>{e(p["nombre"])}<div class="sub">{", ".join("#" + str(f["id"]) for f in p["azure"])}</div></td>'
        f'<td>{chip(p["ultima"]["resultado"] if p["ultima"] else "No ejecutado")}</td>'
        f'<td class="num">{p["ultima"]["duracion"] if p["ultima"] else ""}</td>'
        f'<td>{p["ultima"]["fecha"].strftime("%Y-%m-%d %H:%M:%S") if p["ultima"] else ""}<div class="sub">{len(p["historial"])} ejecución(es)</div></td>'
        f'<td class="msg">{e(p["ultima"]["mensaje"][:300]) if p["ultima"] else ""}'
        f'{"".join(f"<span class=hz>{e(h)}</span>" for h in p["hallazgos"])}</td>'
        f'<td><a href="Reportes_Excel/Por_prueba/{e(p["nombre"])}.xlsx">xlsx</a></td></tr>'
        for p in pruebas)

    hus = {}
    for f in filas:
        for h in f['hu']:
            hus.setdefault((h['id'], h['titulo']), []).append(f['resultado'])
    filas_hu = ''.join(
        f'<tr><td class="num"><a href="{AZURE}{hid}" target="_blank" rel="noopener">{hid}</a></td><td>{e(htit)}</td><td class="num">{len(rs)}</td>'
        f'<td class="num">{rs.count("Pasó")}</td><td class="num">{rs.count("Falló") + rs.count("Error")}</td>'
        f'<td class="num">{rs.count("No ejecutado") + rs.count("No automatizable")}</td>'
        f'<td>{barra({r: rs.count(r) for r in RESULTADOS}, len(rs), "mini")}</td></tr>'
        for (hid, htit), rs in sorted(hus.items(), key=lambda x: x[0][1]))

    filas_det = ''.join(
        f'<tr data-res="{e(f["resultado"])}" data-tipo="{e(f["tipo"])}"><td class="num"><a href="{AZURE}{f["id"]}" target="_blank" rel="noopener">{f["id"]}</a></td>'
        f'<td>{e(f["titulo"])}<div class="sub">{e(f["hu"][0]["titulo"])} · {e(f["responsable"])}</div></td>'
        f'<td>{e(f["tipo"])}<div class="sub">{e(f["katalon"] or "—")}{" · " + e(f["suite"]) if f["suite"] else ""}</div></td>'
        f'<td>{chip(f["resultado"])}</td><td class="num">{f["duracion"] if f["duracion"] is not None else ""}</td>'
        f'<td>{e(f["fecha"])}<div class="sub">{e(f["fuente"])}</div></td>'
        f'<td class="msg">{e(f["mensaje"][:300])}{"<br>" if f["mensaje"] and f["nota"] else ""}{e(f["nota"])}'
        f'{"".join(f"<span class=hz>{e(h)}</span>" for h in f["hallazgos"].split(", ") if h)}</td></tr>'
        for f in filas)

    descargas = f'<a href="Reportes_Excel/{e(excel.name)}">Consolidado · {e(excel.name)}</a>' + ''.join(
        f'<a href="Reportes_Excel/Por_prueba/{e(p["nombre"])}.xlsx">{e(p["caso"])}</a>' for p in pruebas) + \
        '<a href="Documentacion/Informe_Ejecucion_Pruebas.md">Informe .md</a><a href="Documentacion/Inventario_Casos_Azure.md">Inventario Azure</a>'

    pagina = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Resultados de pruebas ITMentorSoft</title>
<style>
{CSS}
</style></head>
<body><main>
<h1>Resultados de pruebas · ITMentorSoft</h1>
<div class="meta">Generado {datetime.now():%Y-%m-%d %H:%M} · Aplicación: <a href="{APP}">{APP}</a> · Casos: {e(datos['fuente'])} ({e(datos['fecha_lectura'])}) · Katalon Studio 11.6, Chrome · Ejecuciones de Katalon leídas: {len(ejecuciones)}</div>
<section class="cards">
{''.join(f'<div class="card"><b>{v}</b><span>{e(t)}</span></div>' for v, t in tarjetas)}
</section>

<div class="panel">{barra(conteo, len(filas))}<ul class="legend">{''.join(f'<li><span class="dot {CLASE[r]}"></span>{r} <b>{conteo[r]}</b></li>' for r in RESULTADOS)}</ul>
<div class="sub" style="margin-top:8px">Resultado de los {len(filas)} casos de Azure según la última ejecución de la prueba Katalon que los cubre.</div></div>

<h2 id="hallazgos">Hallazgos de las pruebas ({len(hz['hallazgos'])}, {defectos} defectos)</h2>
<div class="meta" style="margin-bottom:10px">{e(hz['metodo'])} ({e(hz['fecha'])})</div>
<div class="hzgrid">{hallazgos}</div>

<h2>Pruebas automatizadas en Katalon</h2>
<div class="panel"><table><thead><tr><th class="num">Caso</th><th>Test case Katalon / Azure</th><th>Resultado</th><th class="num">Duración (s)</th><th>Ejecución</th><th>Mensaje / hallazgos</th><th>Reporte</th></tr></thead>
<tbody>{filas_kat}</tbody></table></div>

<h2>Por historia de usuario</h2>
<div class="panel"><table><thead><tr><th class="num">HU</th><th>Historia de usuario</th><th class="num">Casos</th><th class="num">Pasó</th><th class="num">Falló</th><th class="num">Sin ejecutar</th><th></th></tr></thead>
<tbody>{filas_hu}</tbody></table></div>

<h2>Detalle por test case de Azure</h2>
<div class="panel">
<div class="filters"><input id="q" type="search" placeholder="Buscar por ID, título, HU o caso Katalon…" aria-label="Buscar">
<select id="fr" aria-label="Resultado"><option value="">Todos los resultados</option>{''.join(f'<option>{r}</option>' for r in RESULTADOS)}</select>
<select id="ft" aria-label="Tipo"><option value="">Todos los tipos</option>{''.join(f'<option>{t}</option>' for t in TIPOS)}</select></div>
<table id="t"><thead><tr><th class="num">ID</th><th>Test case / HU</th><th>Tipo</th><th>Resultado</th><th class="num">Duración (s)</th><th>Ejecución</th><th>Mensaje / nota</th></tr></thead>
<tbody>{filas_det}</tbody></table></div>

<h2>Reportes descargables</h2>
<div class="downloads">{descargas}</div>
</main>
<script>
const q=document.getElementById('q'),fr=document.getElementById('fr'),ft=document.getElementById('ft');
function filtrar(){{const s=q.value.toLowerCase();document.querySelectorAll('#t tbody tr').forEach(tr=>{{
tr.style.display=(!s||tr.innerText.toLowerCase().includes(s))&&(!fr.value||tr.dataset.res===fr.value)&&(!ft.value||tr.dataset.tipo===ft.value)?'':'none';}});}}
[q,fr,ft].forEach(x=>x.addEventListener('input',filtrar));
</script></body></html>
"""
    destino.write_text(pagina, encoding='utf-8', newline='\n')


def siguiente_excel():
    n = 1
    while (AQUI / f'Reporte_Resultados_{n}.xlsx').exists():
        n += 1
    return AQUI / f'Reporte_Resultados_{n}.xlsx'


def main():
    datos, hz, filas, pruebas, ejecuciones = preparar()
    excel = siguiente_excel()
    excel_consolidado(datos, hz, filas, pruebas, ejecuciones, excel)
    (AQUI / 'Por_prueba').mkdir(exist_ok=True)
    for p in pruebas:
        excel_prueba(p, AQUI / 'Por_prueba' / f"{p['nombre']}.xlsx")
    html_reporte(datos, hz, filas, pruebas, ejecuciones, excel, PROYECTO / 'index.html')
    conteo = {r: sum(f['resultado'] == r for f in filas) for r in RESULTADOS}
    print(f'{excel.name}: {len(filas)} casos de Azure {conteo}; {len(pruebas)} pruebas; {len(ejecuciones)} ejecuciones; index.html')


if __name__ == '__main__':
    main()
